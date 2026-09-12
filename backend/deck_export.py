"""
PPTX export, report markdown, and slide rendering helpers extracted from deck_service.
"""

from __future__ import annotations

import io
from typing import Any, cast

from backend.deck_common import (
    DECK_THEME_PALETTES,
    _clean_text,
    _normalize_chart_datasets,
    _normalize_chart_labels,
    _truncate,
    ensure_deckable_chat,
    normalize_deck_theme,
)
from backend.deck_models import DeckQualityState, DeckSlide, DeckSpec


# ---------------------------------------------------------------------------
# Report markdown (standalone)
# ---------------------------------------------------------------------------
def build_report_markdown(messages: list[Any], title: str) -> str:
    qa_pairs = ensure_deckable_chat(messages)
    lines = [
        "---",
        "theme: default",
        f"title: {title}",
        "class: text-center",
        "---",
        "",
        f"# {title}",
        "",
        "AI 对话报告",
        "",
    ]
    for index, (question, answer) in enumerate(qa_pairs, start=1):
        lines.append("---")
        lines.append("")
        lines.append(f"## 涓婚 {index}: {_truncate(question, 72)}")
        lines.append("")
        lines.append(answer)
        lines.append("")
    return "\n".join(lines)


def _slide_text(slide: DeckSlide) -> str:
    lines: list[str] = []
    for block in slide.blocks:
        if block.kind == "bullet_list":
            for item in block.content.get("items", []):
                value = _clean_text(item)
                if value:
                    lines.append(f"• {value}")
        elif block.kind == "chart":
            title = _clean_text(block.content.get("title")) or "图表"
            chart_type = _clean_text(block.content.get("chart_type")).lower()
            lines.append(f"[图表] {title} ({chart_type or 'chart'})")
        else:
            text = _clean_text(block.content.get("text", ""))
            if text:
                lines.append(text)
    chart_specs = _extract_chart_specs(slide)
    if chart_specs:
        lines.append("")
        lines.append("图表")
        for chart in chart_specs:
            chart_line = f"- {chart['title']} [{chart['chart_type']}]"
            if chart["description"]:
                chart_line += f": {chart['description']}"
            lines.append(chart_line)

    if slide.evidence_refs:
        lines.append("")
        lines.append(
            "Sources: " + ", ".join(ref.source_title for ref in slide.evidence_refs[:3])
        )
    return "\n".join(lines).strip()


def _quality_state_color_key(quality_state: DeckQualityState) -> str:
    if quality_state == "supported":
        return "success"
    if quality_state == "manual":
        return "danger"
    return "warning"


def _is_slide_manually_confirmed(slide: DeckSlide) -> bool:
    return (
        slide.quality_state != "supported" and slide.status.review_state == "confirmed"
    )


def _slide_quality_label(slide: DeckSlide) -> str:
    if _is_slide_manually_confirmed(slide):
        return "已人工确认"
    return _quality_state_label(slide.quality_state)


def _slide_quality_color_key(slide: DeckSlide) -> str:
    if _is_slide_manually_confirmed(slide):
        return "success"
    return _quality_state_color_key(slide.quality_state)


def _extract_chart_specs(slide: DeckSlide) -> list[dict[str, Any]]:
    specs: list[dict[str, Any]] = []
    for block in slide.blocks:
        if block.kind != "chart":
            continue
        chart_type = _clean_text(
            block.content.get("chart_type") or block.content.get("type")
        ).lower()
        if chart_type not in {"bar", "line", "pie"}:
            continue

        labels = _normalize_chart_labels(block.content.get("labels"))
        datasets = _normalize_chart_datasets(block.content.get("datasets"), len(labels))
        if not datasets:
            continue

        if not labels:
            labels = [f"类别 {index + 1}" for index in range(len(datasets[0]["data"]))]
            datasets = _normalize_chart_datasets(
                block.content.get("datasets"), len(labels)
            )
            if not datasets:
                continue

        specs.append(
            {
                "title": _clean_text(block.content.get("title")) or "数据图表",
                "description": _clean_text(block.content.get("description")),
                "chart_type": chart_type,
                "labels": labels[:12],
                "datasets": datasets[:4],
            }
        )
    return specs


def _extract_slide_sections(slide: DeckSlide) -> list[dict[str, Any]]:
    sections: list[dict[str, Any]] = []
    for block in slide.blocks:
        if block.kind == "bullet_list":
            items = [
                _clean_text(item)
                for item in block.content.get("items", [])
                if _clean_text(item)
            ]
            if items:
                sections.append(
                    {
                        "kind": "bullet_list",
                        "role": block.role,
                        "items": items,
                    }
                )
        else:
            text = _clean_text(block.content.get("text", ""))
            if text:
                sections.append(
                    {
                        "kind": "paragraph",
                        "role": block.role,
                        "text": text,
                    }
                )
    return sections


def _quality_state_label(quality_state: DeckQualityState) -> str:
    if quality_state == "supported":
        return "证据充分"
    if quality_state == "manual":
        return "需人工确认"
    return "证据偏弱"


def _split_evenly(items: list[str], column_count: int = 2) -> list[list[str]]:
    if column_count <= 1:
        return [items]
    if not items:
        return [[] for _ in range(column_count)]

    per_column = (len(items) + column_count - 1) // column_count
    return [
        items[index * per_column : (index + 1) * per_column]
        for index in range(column_count)
    ]


def _slide_citation_marker(slide: DeckSlide, *, limit: int = 3) -> str:
    # Keep PPTX body citations compact and aligned with the evidence side panel.
    ref_count = len(slide.evidence_refs or [])
    if ref_count <= 0:
        return ""

    marker_count = min(ref_count, max(1, limit))
    marker = "".join(f"[{index}]" for index in range(1, marker_count + 1))
    return f"{marker}+" if ref_count > marker_count else marker


def _append_citation_marker(
    text: str,
    marker: str,
    *,
    limit: int | None = None,
) -> str:
    cleaned = _clean_text(text)
    if not cleaned:
        return ""
    if not marker or cleaned.endswith(marker):
        return _truncate(cleaned, limit) if limit else cleaned
    if limit:
        suffix = f" {marker}"
        cleaned = _truncate(cleaned, max(1, limit - len(suffix)))
        return f"{cleaned}{suffix}"
    return f"{cleaned} {marker}"


def _compose_export_notes(slide: DeckSlide) -> str:
    lines: list[str] = [slide.title.strip() or "Untitled Slide"]

    if slide.subtitle.strip():
        lines.append(slide.subtitle.strip())

    lines.append(f"质量状态: {_quality_state_label(slide.quality_state)}")

    lines[-1] = f"质量状态: {_slide_quality_label(slide)}"
    if slide.speaker_notes.strip():
        lines.extend(["", "讲述备注", slide.speaker_notes.strip()])
        lines.extend(["", "演讲备注", slide.speaker_notes.strip()])
    sections = _extract_slide_sections(slide)
    if sections:
        lines.append("")
        lines.append("页面内容")
        for section in sections:
            role = _clean_text(section.get("role", ""))
            if role and role not in {"main_points", "summary", "outline", "sources"}:
                lines.append(f"[{role}]")
            if section["kind"] == "bullet_list":
                lines.extend(f"- {item}" for item in section["items"])
            else:
                lines.append(section["text"])

    if slide.evidence_refs:
        lines.append("")
        lines.append("证据来源")
        for ref in slide.evidence_refs:
            confidence = (
                f" ({round(ref.confidence * 100)}%)"
                if ref.confidence and ref.confidence > 0
                else ""
            )
            snippet = _truncate(_clean_text(ref.snippet), 220)
            if snippet:
                lines.append(f"- {ref.source_title}{confidence}: {snippet}")
            else:
                lines.append(f"- {ref.source_title}{confidence}")

    return "\n".join(line for line in lines if line is not None).strip()


def export_deck_to_pptx(deck: DeckSpec) -> bytes:
    try:
        from pptx import Presentation
        from pptx.chart.data import CategoryChartData
        from pptx.dml.color import RGBColor
        from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
        from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
        from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
        from pptx.util import Emu, Pt
    except ImportError as exc:
        raise RuntimeError(
            "python-pptx is not installed. Install requirements.txt and retry."
        ) from exc

    theme = DECK_THEME_PALETTES[normalize_deck_theme(deck.meta.theme)]

    def rgb(color_key: str) -> RGBColor:
        return cast(RGBColor, RGBColor.from_string(theme.get(color_key, color_key)))

    def set_slide_background(ppt_slide, color_key: str = "bg") -> None:
        fill = ppt_slide.background.fill
        fill.solid()
        fill.fore_color.rgb = rgb(color_key)

    def add_textbox(
        ppt_slide,
        x: float,
        y: float,
        w: float,
        h: float,
        *,
        fill_color: str | None = None,
        line_color: str | None = None,
        margins: tuple[float, float, float, float] = (0.08, 0.05, 0.08, 0.05),
        vertical_anchor=MSO_ANCHOR.TOP,
    ):
        from pptx.util import Inches

        shape = ppt_slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        text_frame = shape.text_frame
        text_frame.clear()
        text_frame.word_wrap = True
        text_frame.auto_size = MSO_AUTO_SIZE.NONE
        text_frame.vertical_anchor = vertical_anchor
        text_frame.margin_left = Inches(margins[0])
        text_frame.margin_top = Inches(margins[1])
        text_frame.margin_right = Inches(margins[2])
        text_frame.margin_bottom = Inches(margins[3])

        if fill_color:
            shape.fill.solid()
            shape.fill.fore_color.rgb = rgb(fill_color)
        else:
            shape.fill.background()

        if line_color:
            shape.line.color.rgb = rgb(line_color)
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()

        return shape, text_frame

    def set_text_frame_content(text_frame, paragraphs: list[dict[str, Any]]) -> None:
        text_frame.clear()
        if not paragraphs:
            return

        for index, spec in enumerate(paragraphs):
            paragraph = (
                text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
            )
            paragraph.text = spec.get("text", "")
            paragraph.alignment = spec.get("align", PP_ALIGN.LEFT)
            paragraph.space_after = Pt(spec.get("space_after", 6))
            paragraph.space_before = Pt(spec.get("space_before", 0))
            paragraph.line_spacing = spec.get("line_spacing", 1.15)
            font = paragraph.font
            font.name = spec.get("font_name", "Microsoft YaHei")
            font.size = Pt(spec.get("font_size", 18))
            font.bold = spec.get("bold", False)
            font.italic = spec.get("italic", False)
            font.color.rgb = rgb(spec.get("color", "body"))

    def add_badge(
        ppt_slide,
        text: str,
        x: float,
        y: float,
        w: float,
        h: float,
        *,
        fill_color: str,
        text_color: str = "surface",
    ) -> None:
        from pptx.util import Inches

        shape = ppt_slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            Inches(x),
            Inches(y),
            Inches(w),
            Inches(h),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill_color)
        shape.line.fill.background()
        text_frame = shape.text_frame
        text_frame.word_wrap = False
        text_frame.auto_size = MSO_AUTO_SIZE.NONE
        text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        text_frame.margin_left = Pt(4)
        text_frame.margin_right = Pt(4)
        set_text_frame_content(
            text_frame,
            [
                {
                    "text": text,
                    "font_size": 10,
                    "bold": True,
                    "color": text_color,
                    "align": PP_ALIGN.CENTER,
                    "space_after": 0,
                }
            ],
        )

    def quality_color(slide: DeckSlide) -> str:
        return _slide_quality_color_key(slide)

    def add_chart_panel(
        ppt_slide,
        chart: dict[str, Any],
        x: float,
        y: float,
        w: float,
        h: float,
    ) -> None:
        from pptx.util import Inches

        panel = ppt_slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            Inches(x),
            Inches(y),
            Inches(w),
            Inches(h),
        )
        panel.fill.solid()
        panel.fill.fore_color.rgb = rgb("surface_alt")
        panel.line.color.rgb = rgb("border")
        panel.line.width = Pt(1)

        title_height = 0.36
        description_height = 0.0
        description = _clean_text(chart.get("description"))
        if description:
            description_height = 0.36

        _, title_box = add_textbox(
            ppt_slide,
            x + 0.18,
            y + 0.12,
            w - 0.36,
            title_height,
            margins=(0.0, 0.0, 0.0, 0.0),
        )
        set_text_frame_content(
            title_box,
            [
                {
                    "text": chart["title"],
                    "font_size": 11.5,
                    "bold": True,
                    "color": "title",
                    "space_after": 0,
                }
            ],
        )

        if description:
            _, desc_box = add_textbox(
                ppt_slide,
                x + 0.18,
                y + 0.46,
                w - 0.36,
                description_height,
                margins=(0.0, 0.0, 0.0, 0.0),
            )
            set_text_frame_content(
                desc_box,
                [
                    {
                        "text": _truncate(description, 110),
                        "font_size": 9.5,
                        "color": "muted",
                        "space_after": 0,
                    }
                ],
            )

        chart_box_y = y + 0.56 + description_height
        chart_box_h = max(1.2, h - (chart_box_y - y) - 0.18)
        chart_data = CategoryChartData()
        chart_data.categories = chart["labels"]
        datasets = chart["datasets"]
        if chart["chart_type"] == "pie":
            datasets = datasets[:1]
        for dataset in datasets:
            chart_data.add_series(dataset["label"], tuple(dataset["data"]))

        chart_type_map = {
            "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
            "line": XL_CHART_TYPE.LINE_MARKERS,
            "pie": XL_CHART_TYPE.PIE,
        }
        graphic_frame = ppt_slide.shapes.add_chart(
            chart_type_map[chart["chart_type"]],
            Inches(x + 0.16),
            Inches(chart_box_y),
            Inches(w - 0.32),
            Inches(chart_box_h),
            chart_data,
        )
        ppt_chart = graphic_frame.chart
        ppt_chart.has_title = False
        ppt_chart.has_legend = len(datasets) > 1 or chart["chart_type"] == "pie"
        if ppt_chart.has_legend:
            ppt_chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            ppt_chart.legend.font.size = Pt(9)

        palette = ["accent", "success", "warning", "danger"]
        try:
            if chart["chart_type"] == "pie":
                series = ppt_chart.series[0]
                for point_index, point in enumerate(series.points):
                    color_key = palette[point_index % len(palette)]
                    point.format.fill.solid()
                    point.format.fill.fore_color.rgb = rgb(color_key)
                    point.format.line.color.rgb = rgb("surface")
            else:
                for series_index, series in enumerate(ppt_chart.series):
                    color_key = palette[series_index % len(palette)]
                    series.format.fill.solid()
                    series.format.fill.fore_color.rgb = rgb(color_key)
                    series.format.line.color.rgb = rgb(color_key)
        except Exception:
            pass

    def add_footer(
        ppt_slide,
        slide_index: int,
        total_slides: int,
        slide: DeckSlide,
    ) -> None:
        from pptx.util import Inches

        divider = ppt_slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE,
            Inches(0.75),
            Inches(6.78),
            Inches(11.85),
            Inches(0.02),
        )
        divider.fill.solid()
        divider.fill.fore_color.rgb = rgb("border")
        divider.line.fill.background()

        _, left_footer = add_textbox(
            ppt_slide,
            0.78,
            6.84,
            7.6,
            0.28,
            margins=(0.0, 0.0, 0.0, 0.0),
        )
        set_text_frame_content(
            left_footer,
            [
                {
                    "text": f"{deck.meta.source_mode} | {slide.type} | {slide.layout}",
                    "font_size": 9.5,
                    "color": "muted",
                    "space_after": 0,
                }
            ],
        )

        _, right_footer = add_textbox(
            ppt_slide,
            10.55,
            6.82,
            1.75,
            0.3,
            margins=(0.0, 0.0, 0.0, 0.0),
        )
        set_text_frame_content(
            right_footer,
            [
                {
                    "text": f"{slide_index + 1}/{total_slides}",
                    "font_size": 10,
                    "bold": True,
                    "color": "muted",
                    "align": PP_ALIGN.RIGHT,
                    "space_after": 0,
                }
            ],
        )

    def add_notes(ppt_slide, slide: DeckSlide) -> None:
        notes_text = _compose_export_notes(slide)
        if notes_text:
            ppt_slide.notes_slide.notes_text_frame.text = notes_text

    def format_date_label(raw_value: str) -> str:
        raw = _clean_text(raw_value)
        if not raw:
            return "鏈煡鏃堕棿"
        if "T" in raw:
            return raw.split("T", 1)[0]
        return raw[:10]

    def collect_bullet_items(slide: DeckSlide) -> list[str]:
        items: list[str] = []
        for section in _extract_slide_sections(slide):
            if section["kind"] == "bullet_list":
                items.extend(section["items"])
        return items

    def collect_paragraph_texts(slide: DeckSlide) -> list[str]:
        items: list[str] = []
        for section in _extract_slide_sections(slide):
            if section["kind"] == "paragraph":
                items.append(section["text"])
        return items

    def render_cover_slide(ppt_slide, slide_index: int, slide: DeckSlide) -> None:
        from pptx.util import Inches

        set_slide_background(ppt_slide, "surface")
        top_bar = ppt_slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE,
            Inches(0),
            Inches(0),
            Inches(13.333),
            Inches(0.22),
        )
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = rgb("accent")
        top_bar.line.fill.background()

        accent_panel = ppt_slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            Inches(10.85),
            Inches(0.45),
            Inches(1.85),
            Inches(5.55),
        )
        accent_panel.fill.solid()
        accent_panel.fill.fore_color.rgb = rgb("surface_alt")
        accent_panel.line.fill.background()

        core_message = ""
        for section in _extract_slide_sections(slide):
            if section["kind"] == "paragraph":
                core_message = section["text"]
                break
            if section["kind"] == "bullet_list":
                core_message = "; ".join(section["items"][:2])
                break

        _, eyebrow = add_textbox(
            ppt_slide,
            0.9,
            0.55,
            3.4,
            0.35,
            margins=(0.0, 0.0, 0.0, 0.0),
        )
        set_text_frame_content(
            eyebrow,
            [
                {
                    "text": "AI Knowledge Base Deck",
                    "font_size": 11,
                    "bold": True,
                    "color": "accent",
                    "space_after": 0,
                }
            ],
        )

        _, title_box = add_textbox(
            ppt_slide,
            0.9,
            1.0,
            10.9,
            1.4,
            margins=(0.0, 0.0, 0.0, 0.0),
        )
        set_text_frame_content(
            title_box,
            [
                {
                    "text": slide.title or deck.meta.title,
                    "font_size": 26,
                    "bold": True,
                    "color": "title",
                    "line_spacing": 1.05,
                    "space_after": 0,
                }
            ],
        )

        _, subtitle_box = add_textbox(
            ppt_slide,
            0.92,
            2.35,
            9.1,
            0.75,
            margins=(0.0, 0.0, 0.0, 0.0),
        )
        set_text_frame_content(
            subtitle_box,
            [
                {
                    "text": slide.subtitle
                    or deck.meta.subtitle
                    or "AI 鑷姩鐢熸垚鐨勭粨鏋勫寲姹囨姤鑽夌",
                    "font_size": 16,
                    "color": "muted",
                    "line_spacing": 1.15,
                    "space_after": 0,
                }
            ],
        )

        if core_message:
            _, message_box = add_textbox(
                ppt_slide,
                0.9,
                3.25,
                6.7,
                1.45,
                fill_color="surface_alt",
                line_color="accent_soft",
                margins=(0.18, 0.12, 0.18, 0.12),
            )
            set_text_frame_content(
                message_box,
                [
                    {
                        "text": "鏍稿績缁撹",
                        "font_size": 11,
                        "bold": True,
                        "color": "accent",
                        "space_after": 6,
                    },
                    {
                        "text": core_message,
                        "font_size": 15,
                        "color": "body",
                        "line_spacing": 1.18,
                        "space_after": 0,
                    },
                ],
            )

        _, meta_box = add_textbox(
            ppt_slide,
            8.4,
            3.28,
            3.9,
            2.05,
            fill_color="bg",
            line_color="border",
            margins=(0.16, 0.12, 0.16, 0.12),
        )
        set_text_frame_content(
            meta_box,
            [
                {
                    "text": "瀵煎嚭鎽樿",
                    "font_size": 11,
                    "bold": True,
                    "color": "title",
                    "space_after": 8,
                },
                {
                    "text": f"来源模式: {deck.meta.source_mode}",
                    "font_size": 11,
                    "color": "body",
                    "space_after": 4,
                },
                {
                    "text": f"受众: {deck.meta.audience}",
                    "font_size": 11,
                    "color": "body",
                    "space_after": 4,
                },
                {
                    "text": f"页数: {deck.generation.actual_slide_count}",
                    "font_size": 11,
                    "color": "body",
                    "space_after": 4,
                },
                {
                    "text": f"生成面板: {deck.meta.generator_panel_id}",
                    "font_size": 11,
                    "color": "body",
                    "space_after": 4,
                },
                {
                    "text": f"日期: {format_date_label(deck.meta.created_at)}",
                    "font_size": 11,
                    "color": "body",
                    "space_after": 0,
                },
            ],
        )

        add_badge(
            ppt_slide,
            _slide_quality_label(slide),
            10.55,
            0.62,
            1.72,
            0.38,
            fill_color=quality_color(slide),
        )
        add_footer(ppt_slide, slide_index, len(deck.slides), slide)
        add_notes(ppt_slide, slide)

    def render_outline_slide(ppt_slide, slide_index: int, slide: DeckSlide) -> None:
        set_slide_background(ppt_slide)
        add_badge(
            ppt_slide,
            _slide_quality_label(slide),
            10.55,
            0.55,
            1.72,
            0.38,
            fill_color=quality_color(slide),
        )

        _, title_box = add_textbox(
            ppt_slide,
            0.8,
            0.65,
            8.9,
            0.65,
            margins=(0.0, 0.0, 0.0, 0.0),
        )
        set_text_frame_content(
            title_box,
            [
                {
                    "text": slide.title,
                    "font_size": 24,
                    "bold": True,
                    "color": "title",
                    "space_after": 0,
                }
            ],
        )

        if slide.subtitle:
            _, subtitle_box = add_textbox(
                ppt_slide,
                0.82,
                1.18,
                8.8,
                0.45,
                margins=(0.0, 0.0, 0.0, 0.0),
            )
            set_text_frame_content(
                subtitle_box,
                [
                    {
                        "text": slide.subtitle,
                        "font_size": 13,
                        "color": "muted",
                        "space_after": 0,
                    }
                ],
            )

        outline_items: list[str] = []
        for section in _extract_slide_sections(slide):
            if section["kind"] == "bullet_list":
                outline_items.extend(section["items"])

        visible_items = outline_items[:6]
        card_width = 3.35
        card_height = 1.08
        gap_x = 0.25
        gap_y = 0.22
        start_x = 0.82
        start_y = 1.86

        if not visible_items:
            _, body_box = add_textbox(
                ppt_slide,
                0.82,
                1.8,
                7.1,
                4.55,
                fill_color="surface",
                line_color="border",
                margins=(0.18, 0.14, 0.18, 0.12),
            )
            set_text_frame_content(
                body_box,
                [
                    {
                        "text": "No outline content was generated.",
                        "font_size": 16,
                        "color": "muted",
                        "space_after": 0,
                    }
                ],
            )
        else:
            for item_index, item in enumerate(visible_items, start=1):
                row = (item_index - 1) // 2
                col = (item_index - 1) % 2
                x = start_x + col * (card_width + gap_x)
                y = start_y + row * (card_height + gap_y)
                fill_color = "surface_alt" if item_index == 1 else "surface"
                line_color = "accent_soft" if item_index == 1 else "border"
                _, card_box = add_textbox(
                    ppt_slide,
                    x,
                    y,
                    card_width,
                    card_height,
                    fill_color=fill_color,
                    line_color=line_color,
                    margins=(0.18, 0.14, 0.18, 0.12),
                )
                set_text_frame_content(
                    card_box,
                    [
                        {
                            "text": f"Part {item_index}",
                            "font_size": 10.5,
                            "bold": True,
                            "color": "accent",
                            "space_after": 8,
                        },
                        {
                            "text": item,
                            "font_size": 15.5,
                            "bold": item_index == 1,
                            "color": "body",
                            "line_spacing": 1.12,
                            "space_after": 0,
                        },
                    ],
                )

        _, info_box = add_textbox(
            ppt_slide,
            8.3,
            1.8,
            4.05,
            4.55,
            fill_color="surface_alt",
            line_color="border",
            margins=(0.16, 0.14, 0.16, 0.12),
        )
        warnings_text = (
            "; ".join(w.message for w in deck.generation.warnings[:2])
            if deck.generation.warnings
            else "DeckSpec structured export is being used."
        )
        set_text_frame_content(
            info_box,
            [
                {
                    "text": "导出信息",
                    "font_size": 12,
                    "bold": True,
                    "color": "title",
                    "space_after": 10,
                },
                {
                    "text": f"受众: {deck.meta.audience}",
                    "font_size": 11,
                    "color": "body",
                    "space_after": 6,
                },
                {
                    "text": f"鐩殑: {deck.meta.purpose}",
                    "font_size": 11,
                    "color": "body",
                    "space_after": 6,
                },
                {
                    "text": f"来源模式: {deck.meta.source_mode}",
                    "font_size": 11,
                    "color": "body",
                    "space_after": 6,
                },
                {
                    "text": f"风险提示: {warnings_text}",
                    "font_size": 10.5,
                    "color": "muted",
                    "line_spacing": 1.2,
                    "space_after": 0,
                },
            ],
        )

        add_footer(ppt_slide, slide_index, len(deck.slides), slide)
        add_notes(ppt_slide, slide)

    def render_content_slide(ppt_slide, slide_index: int, slide: DeckSlide) -> None:
        set_slide_background(ppt_slide, "surface")
        add_badge(
            ppt_slide,
            _slide_quality_label(slide),
            10.55,
            0.48,
            1.72,
            0.38,
            fill_color=quality_color(slide),
        )

        from pptx.util import Inches

        accent_rule = ppt_slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE,
            Inches(0.8),
            Inches(0.48),
            Inches(0.09),
            Inches(0.82),
        )
        accent_rule.fill.solid()
        accent_rule.fill.fore_color.rgb = rgb("accent")
        accent_rule.line.fill.background()

        _, title_box = add_textbox(
            ppt_slide,
            1.0,
            0.56,
            9.1,
            0.65,
            margins=(0.0, 0.0, 0.0, 0.0),
        )
        set_text_frame_content(
            title_box,
            [
                {
                    "text": slide.title,
                    "font_size": 23,
                    "bold": True,
                    "color": "title",
                    "space_after": 0,
                }
            ],
        )

        if slide.subtitle:
            _, subtitle_box = add_textbox(
                ppt_slide,
                1.02,
                1.08,
                8.7,
                0.42,
                margins=(0.0, 0.0, 0.0, 0.0),
            )
            set_text_frame_content(
                subtitle_box,
                [
                    {
                        "text": slide.subtitle,
                        "font_size": 12.5,
                        "color": "muted",
                        "space_after": 0,
                    }
                ],
            )

        sections = _extract_slide_sections(slide)
        chart_specs = _extract_chart_specs(slide)
        primary_chart = chart_specs[0] if chart_specs else None
        citation_marker = _slide_citation_marker(slide)
        bullet_items = collect_bullet_items(slide)
        paragraph_texts = collect_paragraph_texts(slide)
        summary_text = (
            paragraph_texts[0]
            if paragraph_texts
            else _clean_text(slide.subtitle or slide.intent)
        )

        body_x = 0.8
        body_y = 1.65
        body_width = 7.85 if slide.evidence_refs else 11.45
        body_height = 4.95

        if summary_text:
            _, summary_box = add_textbox(
                ppt_slide,
                body_x,
                1.62,
                body_width,
                0.78,
                fill_color="surface_alt",
                line_color="accent_soft",
                margins=(0.18, 0.12, 0.18, 0.1),
            )
            set_text_frame_content(
                summary_box,
                [
                    {
                        "text": "关键摘要",
                        "font_size": 10.5,
                        "bold": True,
                        "color": "accent",
                        "space_after": 5,
                    },
                    {
                        "text": _append_citation_marker(
                            summary_text,
                            citation_marker,
                            limit=120,
                        ),
                        "font_size": 13.5,
                        "color": "body",
                        "line_spacing": 1.15,
                        "space_after": 0,
                    },
                ],
            )
            body_y = 2.55
            body_height = 4.05

        chart_panel: tuple[float, float, float, float] | None = None
        if primary_chart and slide.evidence_refs:
            body_height = 1.55
            chart_panel = (0.8, 4.15, 7.85, 2.25)
        elif primary_chart:
            body_width = 5.15
            chart_panel = (6.2, body_y, 6.05, body_height)

        bullet_count = len(bullet_items)
        max_bullet_length = max((len(item) for item in bullet_items), default=0)
        use_cards = (
            not slide.evidence_refs
            and primary_chart is None
            and 1 < bullet_count <= 4
            and max_bullet_length <= 34
        )
        use_two_columns = (
            not slide.evidence_refs and primary_chart is None and bullet_count >= 5
        )

        if use_cards:
            card_width = 5.55
            card_height = 1.24 if bullet_count <= 2 else 1.08
            gap_x = 0.35
            gap_y = 0.22
            for item_index, item in enumerate(bullet_items[:4], start=1):
                row = (item_index - 1) // 2
                col = (item_index - 1) % 2
                x = body_x + col * (card_width + gap_x)
                y = body_y + row * (card_height + gap_y)
                _, card_box = add_textbox(
                    ppt_slide,
                    x,
                    y,
                    card_width,
                    card_height,
                    fill_color="bg" if item_index % 2 == 1 else "surface_alt",
                    line_color="border",
                    margins=(0.16, 0.12, 0.16, 0.1),
                )
                set_text_frame_content(
                    card_box,
                    [
                        {
                            "text": f"要点 {item_index}",
                            "font_size": 10.5,
                            "bold": True,
                            "color": "accent",
                            "space_after": 7,
                        },
                        {
                            "text": item,
                            "font_size": 16,
                            "color": "body",
                            "line_spacing": 1.15,
                            "space_after": 0,
                        },
                    ],
                )
        elif use_two_columns:
            split_columns = _split_evenly(bullet_items, column_count=2)
            column_width = 5.55
            gap_x = 0.35
            for column_index, items in enumerate(split_columns):
                x = body_x + column_index * (column_width + gap_x)
                _, column_box = add_textbox(
                    ppt_slide,
                    x,
                    body_y,
                    column_width,
                    body_height,
                    fill_color="bg",
                    line_color="border",
                    margins=(0.18, 0.14, 0.18, 0.12),
                )
                column_paragraphs: list[dict[str, Any]] = [
                    {
                        "text": f"关键要点 {column_index + 1}",
                        "font_size": 10.5,
                        "bold": True,
                        "color": "accent",
                        "space_after": 8,
                    }
                ]
                for item in items:
                    column_paragraphs.append(
                        {
                            "text": f"• {item}",
                            "font_size": 14.5,
                            "color": "body",
                            "line_spacing": 1.16,
                            "space_after": 8,
                        }
                    )
                set_text_frame_content(column_box, column_paragraphs)
        else:
            _, body_box = add_textbox(
                ppt_slide,
                body_x,
                body_y,
                body_width,
                body_height,
                fill_color="bg",
                line_color="border",
                margins=(0.18, 0.14, 0.18, 0.12),
            )

            body_font_size = 17 if bullet_count <= 4 else 15
            body_paragraphs: list[dict[str, Any]] = []
            for section in sections:
                role = _clean_text(section.get("role", ""))
                if role and role not in {"main_points", "summary"}:
                    body_paragraphs.append(
                        {
                            "text": role.replace("_", " ").title(),
                            "font_size": 10.5,
                            "bold": True,
                            "color": "accent",
                            "space_after": 6,
                        }
                    )
                if section["kind"] == "bullet_list":
                    for item in section["items"]:
                        display_item = _append_citation_marker(item, citation_marker)
                        body_paragraphs.append(
                            {
                                "text": f"• {display_item}",
                                "font_size": body_font_size,
                                "color": "body",
                                "line_spacing": 1.18,
                                "space_after": 10,
                            }
                        )
                else:
                    display_text = _append_citation_marker(
                        section["text"], citation_marker
                    )
                    body_paragraphs.append(
                        {
                            "text": display_text,
                            "font_size": 15,
                            "color": "body",
                            "line_spacing": 1.2,
                            "space_after": 10,
                        }
                    )
            set_text_frame_content(body_box, body_paragraphs)

        if primary_chart and chart_panel:
            add_chart_panel(
                ppt_slide,
                primary_chart,
                chart_panel[0],
                chart_panel[1],
                chart_panel[2],
                chart_panel[3],
            )

        if slide.evidence_refs:
            _, evidence_box = add_textbox(
                ppt_slide,
                8.95,
                1.62,
                3.55,
                4.98,
                fill_color="surface_alt",
                line_color="border",
                margins=(0.16, 0.14, 0.16, 0.12),
            )
            evidence_paragraphs: list[dict[str, Any]] = [
                {
                    "text": "证据来源",
                    "font_size": 12,
                    "bold": True,
                    "color": "title",
                    "space_after": 8,
                }
            ]
            for ref_index, ref in enumerate(slide.evidence_refs[:3], start=1):
                confidence = (
                    f" ({round(ref.confidence * 100)}%)"
                    if ref.confidence and ref.confidence > 0
                    else ""
                )
                evidence_paragraphs.append(
                    {
                        "text": f"[{ref_index}] {ref.source_title}{confidence}",
                        "font_size": 10.5,
                        "bold": True,
                        "color": "body",
                        "space_after": 4,
                    }
                )
                evidence_paragraphs.append(
                    {
                        "text": _truncate(_clean_text(ref.snippet), 150) or "No excerpt.",
                        "font_size": 9.5,
                        "color": "muted",
                        "line_spacing": 1.2,
                        "space_after": 10,
                    }
                )
            evidence_paragraphs.append(
                {
                    "text": f"状态: {_slide_quality_label(slide)}",
                    "font_size": 10,
                    "bold": True,
                    "color": quality_color(slide),
                    "space_after": 0,
                }
            )
            set_text_frame_content(evidence_box, evidence_paragraphs)

        add_footer(ppt_slide, slide_index, len(deck.slides), slide)
        add_notes(ppt_slide, slide)

    def render_appendix_slide(ppt_slide, slide_index: int, slide: DeckSlide) -> None:
        set_slide_background(ppt_slide)
        add_badge(
            ppt_slide,
            _slide_quality_label(slide),
            10.55,
            0.55,
            1.72,
            0.38,
            fill_color=quality_color(slide),
        )

        _, title_box = add_textbox(
            ppt_slide,
            0.8,
            0.65,
            10.0,
            0.62,
            margins=(0.0, 0.0, 0.0, 0.0),
        )
        set_text_frame_content(
            title_box,
            [
                {
                    "text": slide.title,
                    "font_size": 22,
                    "bold": True,
                    "color": "title",
                    "space_after": 0,
                }
            ],
        )

        if slide.subtitle:
            _, subtitle_box = add_textbox(
                ppt_slide,
                0.82,
                1.14,
                9.4,
                0.4,
                margins=(0.0, 0.0, 0.0, 0.0),
            )
            set_text_frame_content(
                subtitle_box,
                [
                    {
                        "text": slide.subtitle,
                        "font_size": 12.5,
                        "color": "muted",
                        "space_after": 0,
                    }
                ],
            )

        appendix_items: list[str] = []
        for section in _extract_slide_sections(slide):
            if section["kind"] == "bullet_list":
                appendix_items.extend(section["items"])
            else:
                appendix_items.append(section["text"])

        total_items = len(appendix_items)
        _, stat_box = add_textbox(
            ppt_slide,
            9.58,
            0.72,
            2.88,
            0.82,
            fill_color="surface_alt",
            line_color="border",
            margins=(0.14, 0.1, 0.14, 0.08),
        )
        set_text_frame_content(
            stat_box,
            [
                {
                    "text": "鏉ユ簮缁熻",
                    "font_size": 10.5,
                    "bold": True,
                    "color": "accent",
                    "space_after": 4,
                    "align": PP_ALIGN.CENTER,
                },
                {
                    "text": f"共 {total_items} 个来源",
                    "font_size": 13,
                    "bold": True,
                    "color": "title",
                    "align": PP_ALIGN.CENTER,
                    "space_after": 0,
                },
            ],
        )

        if total_items <= 12:
            card_width = 5.55
            card_height = 0.68 if total_items <= 8 else 0.56
            gap_x = 0.3
            gap_y = 0.16
            for item_index, item in enumerate(appendix_items, start=1):
                row = (item_index - 1) // 2
                col = (item_index - 1) % 2
                x = 0.8 + col * (card_width + gap_x)
                y = 1.78 + row * (card_height + gap_y)
                _, item_box = add_textbox(
                    ppt_slide,
                    x,
                    y,
                    card_width,
                    card_height,
                    fill_color="surface" if item_index % 2 else "bg",
                    line_color="border",
                    margins=(0.14, 0.08, 0.14, 0.06),
                )
                set_text_frame_content(
                    item_box,
                    [
                        {
                            "text": f"{item_index:02d}",
                            "font_size": 9.5,
                            "bold": True,
                            "color": "accent",
                            "space_after": 3,
                        },
                        {
                            "text": item,
                            "font_size": 11.5 if total_items <= 8 else 10.5,
                            "color": "body",
                            "space_after": 0,
                            "line_spacing": 1.05,
                        },
                    ],
                )
        else:
            columns = _split_evenly(appendix_items, column_count=2)
            font_size = 11
            for column_index, items in enumerate(columns):
                x = 0.8 + column_index * 6.1
                _, column_box = add_textbox(
                    ppt_slide,
                    x,
                    1.78,
                    5.45,
                    4.9,
                    fill_color="surface",
                    line_color="border",
                    margins=(0.16, 0.12, 0.16, 0.1),
                )
                column_paragraphs = [
                    {
                        "text": f"来源 {column_index + 1}",
                        "font_size": 11,
                        "bold": True,
                        "color": "accent",
                        "space_after": 8,
                    }
                ]
                column_paragraphs.extend(
                    {
                        "text": f"• {item}",
                        "font_size": font_size,
                        "color": "body",
                        "line_spacing": 1.12,
                        "space_after": 6,
                    }
                    for item in items
                )
                set_text_frame_content(column_box, column_paragraphs)

        add_footer(ppt_slide, slide_index, len(deck.slides), slide)
        add_notes(ppt_slide, slide)

    presentation = Presentation()
    presentation.slide_width = Emu(12192000)
    presentation.slide_height = Emu(6858000)
    presentation.core_properties.title = deck.meta.title
    presentation.core_properties.author = deck.meta.author or "system"
    presentation.core_properties.subject = deck.meta.subtitle or deck.meta.purpose
    presentation.core_properties.keywords = "deck,pptx,insightdesk"

    blank_layout = presentation.slide_layouts[6]
    for index, slide in enumerate(deck.slides):
        ppt_slide = presentation.slides.add_slide(blank_layout)
        if slide.type == "cover" or index == 0:
            render_cover_slide(ppt_slide, index, slide)
        elif slide.type == "outline":
            render_outline_slide(ppt_slide, index, slide)
        elif slide.type.startswith("appendix"):
            render_appendix_slide(ppt_slide, index, slide)
        else:
            render_content_slide(ppt_slide, index, slide)

    buffer = io.BytesIO()
    presentation.save(buffer)
    buffer.seek(0)
    return buffer.read()


def build_export_filename(deck: DeckSpec, extension: str) -> str:
    safe_title = "".join(
        char for char in deck.meta.title if char.isalnum() or char in (" ", "-", "_")
    ).strip()[:40]
    return f"{safe_title or 'deck'}.{extension}"

